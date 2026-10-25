"""Fetch the official NCTS phase 6 and ICS2 XML schemas (and the ICS2 code lists) and write them under
backend/data/xsd/{ncts,ics2}/<version>/ with a SHA256SUMS manifest.

Sources (all official DG TAXUD artefacts, see README.md next to each set):
  * NCTS-P6: DDNTA 6.4.1-v1.00 (SfA) Appendix X (XSD 60.4.16), as republished by Swedish Customs (Tullverket) on its
    "Technical specifications for NCTS phase 6" page. Finland runs NCTS-P6 (opt-out) since the P6 deployment window
    ended on 1.9.2025 (DG TAXUD guidance in the CIRCABC ICS2 library, "ICS2 and NCTS-P6" folder).
  * ICS2: "ICS2 3.2.0.0 CTSS EO Package.7z" (Common Technical System Specifications, economic operator package) in the
    public CIRCABC interest group "EU Advance Cargo Information System (ICS2)", folder "ICS2 Common Technical System
    Specifications"; the IE3 XSDs (namespace versions V2 and V3) are in the inner "ICS2 3.2.0.0-EO-TSC.7z".
  * ICS2 code lists: "ICS2-HTI-v3.20-EO.zip" (Harmonised Trader Interface specifications v3.20, 2026-06-30) in the same
    CIRCABC group, folder "ICS2 Common Functional System Specifications"; file "7. ICS2-HTI-CL-(2026-06-30)-v3.20.xlsx".

Usage (needs network; py7zr for the .7z archives: pip install py7zr):
    python backend/data/xsd/fetch_ncts_ics2.py            # writes into a new, empty directory next to this file
    python backend/data/xsd/fetch_ncts_ics2.py --out /tmp/x
Re-running into the committed directories is refused when they already exist (compare with --out instead).
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import sys
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent

NCTS_VERSION = "60.4.16"
NCTS_URL = ("https://www.tullverket.se/download/18.6967fe101964058a61b7a007/1752650865809/"
            "DDNTA_APP_X_6.4.1-v1.00(SfA).7z")
NCTS_PAGE = ("https://www.tullverket.se/foretag/ansokochdeklarera/deklareradigitalt/ediforsystemutvecklare/"
             "editekniskaspecifikationer/sctsnctsfas6.4.72334712192df55a6c0789b3.html")

ICS2_VERSION = "3.2.0.0"
CIRCABC = "https://circabc.europa.eu/rest/download/"
ICS2_CTSS_ID = "5c528a3e-b9b8-4e6f-82ac-272962a92514"          # ICS2 3.2.0.0 CTSS EO Package.7z
ICS2_HTI_ID = "4b6801df-72b8-416a-9b11-14527480d86e"           # ICS2-HTI-v3.20-EO.zip
ICS2_GROUP = ("https://circabc.europa.eu/ui/group/18fb5859-3970-4ac5-b30b-6604977a15a7/library/"
              "16e29912-d290-417e-a45c-b286407e1770")
CODE_LISTS = ["CL013", "CL017", "CL018", "CL094", "CL141", "CL704", "CL718", "CL723", "CL729", "CL735", "CL742",
              "CL744", "CL745", "CL749", "CL750", "CL751", "CL754"]


def get(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (Huolitsija schema fetch)"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()


def sha(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def seven_zip(data: bytes) -> dict[str, bytes]:
    import tempfile

    import py7zr  # pip install py7zr

    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        with py7zr.SevenZipFile(io.BytesIO(data)) as z:
            z.extractall(tmp)
        return {str(p.relative_to(tmp)): p.read_bytes() for p in Path(tmp).rglob("*") if p.is_file()}


def write_set(target: Path, files: dict[str, bytes], meta: dict):
    if target.exists():
        sys.exit(f"{target} exists; use --out <new empty dir> and compare")
    target.mkdir(parents=True)
    lines = []
    for rel, data in sorted(files.items()):
        p = target / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(data)
        lines.append(f"{sha(data)}  {rel}")
    (target / "SHA256SUMS").write_text("\n".join(lines) + "\n")
    (target / "SOURCE.json").write_text(json.dumps(meta, indent=1) + "\n")
    print(f"wrote {len(files)} files to {target}")


def fetch_ncts(out: Path):
    arc = get(NCTS_URL)
    files = {}
    for name, data in seven_zip(arc).items():
        base = name.split("/")[-1]
        if base.lower().endswith(".xsd") and "/RuleMappings/" not in name:
            files[base] = data
    write_set(out / "ncts" / NCTS_VERSION, files, {
        "set": "NCTS phase 6 XML schemas (DDNTA 6.4.1-v1.00 SfA, Appendix X)", "xsd_version": NCTS_VERSION,
        "archive_url": NCTS_URL, "archive_page": NCTS_PAGE, "archive_sha256": sha(arc), "archive_bytes": len(arc),
        "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds")})


def fetch_ics2(out: Path):
    pkg = get(CIRCABC + ICS2_CTSS_ID)
    outer = seven_zip(pkg)
    tsc_name = next(n for n in outer if n.endswith("-EO-TSC.7z"))
    inner = seven_zip(outer[tsc_name])
    files = {}
    for name, data in inner.items():
        parts = name.split("/")
        if "IE3-XSDs" in parts and name.lower().endswith(".xsd"):
            files["/".join(parts[parts.index("IE3-XSDs") + 1:])] = data
        elif parts[-1] == "history.txt":
            files["history.txt"] = data
    hti = get(CIRCABC + ICS2_HTI_ID)
    cl_name = None
    with zipfile.ZipFile(io.BytesIO(hti)) as z:
        for n in z.namelist():
            if n.split("/")[-1].startswith("7. ICS2-HTI-CL-") and n.endswith(".xlsx"):
                cl_name, cl_bytes = n, z.read(n)
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(cl_bytes), read_only=True, data_only=True)
    lists = {}
    for cl in CODE_LISTS:
        if cl not in wb.sheetnames:
            continue
        rows = list(wb[cl].iter_rows(values_only=True))
        head = [str(c or "").strip() for c in rows[0]]
        iv, idsc = head.index("CL Value"), head.index("CL Value Description")
        name = next((r[1] for r in rows[1:] if r[1]), None)
        lists[cl] = {"name": str(name).strip() if name else None,
                     "values": {str(r[iv]).strip(): str(r[idsc] or "").strip() for r in rows[1:] if r[iv] is not None}}
    files["codelists.json"] = (json.dumps({"source": cl_name, "lists": lists}, indent=1, ensure_ascii=False) + "\n").encode()
    write_set(out / "ics2" / ICS2_VERSION, files, {
        "set": "ICS2 IE3 XML schemas (CTSS EO package 3.2.0.0, namespace versions V2 and V3) + HTI v3.20 code lists",
        "package_url": CIRCABC + ICS2_CTSS_ID, "package_name": "ICS2 3.2.0.0 CTSS EO Package.7z",
        "package_sha256": sha(pkg), "package_bytes": len(pkg), "inner_archive": tsc_name,
        "codelists_url": CIRCABC + ICS2_HTI_ID, "codelists_file": cl_name, "codelists_zip_sha256": sha(hti),
        "circabc_group": ICS2_GROUP, "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds")})


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE))
    ap.add_argument("--only", choices=["ncts", "ics2"])
    a = ap.parse_args()
    if a.only != "ics2":
        fetch_ncts(Path(a.out))
    if a.only != "ncts":
        fetch_ics2(Path(a.out))
