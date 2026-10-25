"""Document extraction agent: reads an uploaded document, detects its type and returns structured fields
with a per-field confidence, then runs deterministic checks the model cannot be trusted with (arithmetic,
check digits, code lists). Results land in documents.extraction and `document.extracted` is emitted.

Queue: durable on `documents.extraction.status` (pending -> processing -> done | needs_review | failed).
A worker loop claims one job at a time; kick() wakes it immediately after an upload. A job stuck in
processing for 10 minutes (server restart) is put back to pending, at most 3 attempts.

Model input: text PDFs -> the PDF text (layout-preserving pypdf extraction); scanned PDFs -> page images
rendered with pypdfium2; images -> the image; CSV/XLSX -> the cell text. Everything goes through core.llm so
the provider stays configurable. When the model is unavailable the document is marked needs_review with
the reason: no field is ever invented.
"""
from __future__ import annotations

import asyncio
import csv
import io
import logging
import re
from datetime import date, datetime
from zoneinfo import ZoneInfo

from core.audit import agent_actor, audit
from core.db import db
from core.events import emit
from core.lifecycle import background_loop
from core.llm import LLMUnavailable, llm_available, llm_json
from core.storage import get_file
from core.util import NO_ID, now, now_iso
from services.documents import DOC_TYPES, run_tieout

log = logging.getLogger("extraction")

AGENT = "doc_parser"
HELSINKI = ZoneInfo("Europe/Helsinki")
MAX_PAGES_VISION = 6
MAX_TEXT_CHARS = 60000

# ---------------------------------------------------------------------------------------------
# Field specifications per document type (domain-spec section 8)
# kind: text | date | number | country | currency | incoterm | list | eori | vat | hs
# ---------------------------------------------------------------------------------------------
F = lambda kind="text", required=False: {"kind": kind, "required": required}  # noqa: E731

PARTY = lambda p, req=False: {  # noqa: E731
    f"{p}_name": F("text", req), f"{p}_address": F("text"), f"{p}_country": F("country"),
}

FIELD_SPECS: dict[str, dict] = {
    "commercial_invoice": {
        "fields": {
            "invoice_number": F("text", True), "invoice_date": F("date", True),
            **PARTY("seller", True), "seller_vat_id": F("vat"), "seller_eori": F("eori"),
            "seller_city": F(), "seller_postal_code": F(),
            **PARTY("buyer", True), "buyer_vat_id": F("vat"), "buyer_eori": F("eori"),
            "consignee_name": F(), "incoterm": F("incoterm", True), "incoterm_place": F("text", True),
            "currency": F("currency", True), "subtotal_amount": F("number"), "freight_amount": F("number"),
            "insurance_amount": F("number"), "other_charges_amount": F("number"), "total_amount": F("number", True),
            "transaction_nature": F(), "payment_terms": F(), "country_of_origin": F("country"), "total_packages": F("number"),
            "total_gross_weight_kg": F("number"), "total_net_weight_kg": F("number"),
            "purchase_order_number": F(), "origin_declaration_text": F(),
        },
        "line_items": {"description": F("text", True), "quantity": F("number", True), "unit": F(),
                       "unit_price": F("number"), "amount": F("number", True), "hs_code": F("hs"),
                       "country_of_origin": F("country"), "net_weight_kg": F("number"), "gross_weight_kg": F("number"),
                       "sku": F()},
    },
    "packing_list": {
        "fields": {
            "packing_list_number": F(), "date": F("date"), "invoice_number": F(), **PARTY("shipper"),
            **PARTY("consignee", True), "total_packages": F("number", True), "package_type": F(),
            "total_gross_weight_kg": F("number", True), "total_net_weight_kg": F("number"),
            "total_volume_m3": F("number"), "container_numbers": F("list"), "seal_numbers": F("list"),
            "marks_and_numbers": F(),
        },
        "line_items": {"description": F("text", True), "packages": F("number"), "package_type": F(),
                       "quantity": F("number"), "net_weight_kg": F("number"), "gross_weight_kg": F("number"),
                       "dimensions": F(), "volume_m3": F("number")},
    },
    "bill_of_lading": {
        "fields": {
            "bl_number": F("text", True), "bl_kind": F(), **PARTY("shipper", True), **PARTY("consignee", True),
            "notify_party": F(), "carrier_name": F(), "vessel_flag": F("country"), "vessel_name": F(), "voyage_number": F(),
            "port_of_loading": F("text", True), "port_of_discharge": F("text", True), "place_of_receipt": F(),
            "place_of_delivery": F(), "container_numbers": F("list"), "seal_numbers": F("list"),
            "goods_description": F(), "total_packages": F("number"), "package_type": F(),
            "total_gross_weight_kg": F("number", True), "total_volume_m3": F("number"), "freight_terms": F(),
            "number_of_originals": F("number"), "place_of_issue": F(), "date_of_issue": F("date"),
            "shipped_on_board_date": F("date"),
        },
        "line_items": None,
    },
    "air_waybill": {
        "fields": {
            "awb_number": F("text", True), "house_awb_number": F(), **PARTY("shipper", True),
            **PARTY("consignee", True), "airport_of_departure": F("text", True),
            "airport_of_destination": F("text", True), "flight_numbers": F("list"), "total_packages": F("number"),
            "total_gross_weight_kg": F("number", True), "chargeable_weight_kg": F("number"),
            "nature_of_goods": F(), "handling_information": F(), "date_of_issue": F("date"),
        },
        "line_items": None,
    },
    "cmr": {
        "fields": {
            "cmr_number": F(), **PARTY("sender", True), **PARTY("consignee", True), "carrier_name": F(),
            "place_of_taking_over": F("text", True), "date_of_taking_over": F("date"),
            "place_of_delivery": F("text", True), "goods_description": F(), "total_packages": F("number"),
            "package_type": F(), "total_gross_weight_kg": F("number", True), "vehicle_country": F("country"), "vehicle_plate": F(),
            "trailer_plate": F(), "sender_instructions": F(),
        },
        "line_items": None,
    },
    "certificate_of_origin": {
        "fields": {
            "certificate_number": F("text", True), **PARTY("exporter", True), **PARTY("consignee"),
            "country_of_origin": F("country", True), "goods_description": F(), "invoice_number": F(),
            "issuing_body": F(), "issue_date": F("date"), "total_gross_weight_kg": F("number"),
        },
        "line_items": None,
    },
}
FIELD_SPECS["sea_waybill"] = FIELD_SPECS["bill_of_lading"]
FIELD_SPECS["eur1"] = FIELD_SPECS["certificate_of_origin"]
FIELD_SPECS["cim"] = FIELD_SPECS["cmr"]
FIELD_SPECS["other"] = {"fields": {"document_title": F(), "reference_number": F(), "date": F("date"),
                                   "issuer_name": F(), "summary": F()}, "line_items": None}

INCOTERMS_2020 = {"EXW", "FCA", "CPT", "CIP", "DAP", "DPU", "DDP", "FAS", "FOB", "CFR", "CIF"}
OLD_INCOTERMS = {"DAT": "DPU", "DAF": "DAP", "DDU": "DAP", "DES": "DAP", "DEQ": "DPU", "C&F": "CFR", "CNF": "CFR"}
CURRENCIES = {"EUR", "USD", "GBP", "SEK", "NOK", "DKK", "CHF", "PLN", "CZK", "HUF", "RON", "BGN", "ISK", "TRY",
              "CNY", "JPY", "KRW", "INR", "TWD", "HKD", "SGD", "THB", "VND", "IDR", "MYR", "PHP", "AUD", "NZD", "CAD",
              "MXN", "BRL", "ZAR", "AED", "SAR", "ILS", "EGP", "MAD", "CLP", "COP", "PEN", "ARS", "BDT", "PKR", "LKR",
              "KES", "NGN", "UAH", "RSD", "GEL", "KZT", "QAR", "KWD", "BHD", "OMR", "JOD", "TND", "CNH"}

# ---------------------------------------------------------------------------------------------
# Model input
# ---------------------------------------------------------------------------------------------


def _pdf_text(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        reader.decrypt("")
    parts = []
    for i, page in enumerate(reader.pages):
        try:
            txt = page.extract_text(extraction_mode="layout") or ""
        except Exception:  # noqa: BLE001
            txt = page.extract_text() or ""
        txt = re.sub(r"[ \t]+\n", "\n", txt)
        parts.append(f"--- page {i + 1} ---\n{txt.strip()}")
    return "\n".join(parts), len(reader.pages)


def _render_pages(data: bytes, max_pages: int = MAX_PAGES_VISION) -> list[bytes]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    out = []
    try:
        for i in range(min(len(pdf), max_pages)):
            page = pdf[i]
            w, h = page.get_size()
            scale = min(2.0, 1800 / max(w, h))  # ~1800 px on the long side keeps small print legible
            img = page.render(scale=scale).to_pil()
            buf = io.BytesIO()
            img.convert("RGB").save(buf, format="PNG", optimize=True)
            out.append(buf.getvalue())
    finally:
        pdf.close()
    return out


def _prepare_image(data: bytes, mime: str) -> tuple[bytes, str]:
    from PIL import Image

    img = Image.open(io.BytesIO(data))
    if max(img.size) <= 2200 and len(data) <= 4_000_000 and mime in ("image/png", "image/jpeg", "image/webp"):
        return data, mime
    img.thumbnail((2200, 2200))
    buf = io.BytesIO()
    img.convert("RGB").save(buf, format="JPEG", quality=88)
    return buf.getvalue(), "image/jpeg"


def _sheet_text(data: bytes, mime: str) -> str:
    if mime == "text/csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        return "\n".join("\t".join(c.strip() for c in r) for r in rows[:2000])
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:5]:
        parts.append(f"--- sheet {ws.title} ---")
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r > 2000:
                break
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                parts.append("\t".join(cells).rstrip())
    return "\n".join(parts)


def build_input(data: bytes, mime: str) -> dict:
    """-> {method: text|vision, text, files[], pages}"""
    if mime == "application/pdf":
        text, pages = _pdf_text(data)
        meaningful = len(re.sub(r"\s+|--- page \d+ ---", "", text))
        if meaningful >= 120 * min(pages, 3):
            return {"method": "text", "text": text[:MAX_TEXT_CHARS], "files": [], "pages": pages,
                    "truncated": len(text) > MAX_TEXT_CHARS}
        imgs = _render_pages(data)
        return {"method": "vision", "text": "", "pages": pages, "truncated": pages > MAX_PAGES_VISION,
                "files": [("image/png", b, f"page-{i + 1}.png") for i, b in enumerate(imgs)]}
    if mime.startswith("image/"):
        b, m = _prepare_image(data, mime)
        return {"method": "vision", "text": "", "pages": 1, "files": [(m, b, "image")], "truncated": False}
    text = _sheet_text(data, mime)
    return {"method": "text", "text": text[:MAX_TEXT_CHARS], "files": [], "pages": 1,
            "truncated": len(text) > MAX_TEXT_CHARS}


# ---------------------------------------------------------------------------------------------
# Prompt
# ---------------------------------------------------------------------------------------------
SYSTEM = """You read shipping and customs documents for a Finnish freight forwarder and customs broker.
Return exactly what the document states. Never guess or fill in a value that is not printed on the document:
use null and confidence 0 instead. Normalise only the format:
- dates as YYYY-MM-DD; numbers as plain JSON numbers with a dot decimal (no thousands separators, no units);
- weights in kilograms (convert lb/t when the unit is printed, and say so in warnings);
- countries as ISO 3166-1 alpha-2 codes (e.g. "China" -> "CN"); currencies as ISO 4217 codes;
- incoterm as the 3-letter code, incoterm_place as the named place;
- container numbers as 4 letters + 7 digits without spaces;
- addresses as one line.
confidence is your probability (0..1) that the value is exactly right as printed; lower it for blurry, handwritten,
ambiguous or inferred values."""


def _schema_text(doc_type: str | None) -> str:
    lines = []
    types = [doc_type] if doc_type and doc_type in FIELD_SPECS else [t for t in FIELD_SPECS if t not in ("sea_waybill", "eur1", "cim")]
    for t in types:
        spec = FIELD_SPECS[t]
        f = ", ".join(f"{k}({v['kind']}{', required' if v['required'] else ''})" for k, v in spec["fields"].items())
        lines.append(f"- {t}: fields: {f}")
        if spec.get("line_items"):
            li = ", ".join(spec["line_items"].keys())
            lines.append(f"  line_items keys: {li}")
    return "\n".join(lines)


def build_prompt(declared_type: str | None, inp: dict, filename: str) -> str:
    known = ", ".join(t for t in DOC_TYPES if t not in ("power_of_attorney", "pod", "customs_decision", "dg_declaration"))
    hint = (f"The uploader says this is a {declared_type.replace('_', ' ')}. If it clearly is another type, say so in "
            f"document_type." if declared_type else "The uploader did not say what this document is; detect it.")
    schema = _schema_text(declared_type if declared_type in FIELD_SPECS else None)
    body = (f"Document text (pypdf layout extraction, {inp['pages']} page(s)"
            f"{', truncated' if inp.get('truncated') else ''}):\n<<<\n{inp['text']}\n>>>" if inp["method"] == "text"
            else f"The document is attached as {len(inp['files'])} page image(s)"
                 f"{' (only the first pages)' if inp.get('truncated') else ''}.")
    today = datetime.now(HELSINKI).date().isoformat()
    return f"""Today is {today} (Europe/Helsinki). Judge dates against today: a document dated today or earlier is not
in the future.
File name: {filename}
{hint}
Known document types: {known}.

Field keys per type (use exactly these keys; omit keys you cannot find):
{schema}

{body}

Return JSON:
{{"document_type": "<one known type>", "type_confidence": 0..1, "language": "en|fi|sv|de|zh|...",
  "fields": {{"<key>": {{"value": <string|number|array|null>, "confidence": 0..1}}, ...}},
  "line_items": [{{"<line key>": <value>, ..., "confidence": 0..1}}],
  "warnings": [{{"en": "<note in English>", "fi": "<the same note in Finnish>"}}]}}
warnings: anything a customs broker must know (unreadable parts, stamps over text, conflicting totals, missing pages),
each note written in English and in Finnish with the same facts; keep numbers, codes and names exactly as printed.
Include every line item row (not subtotals). For types without line items return "line_items": []."""


# ---------------------------------------------------------------------------------------------
# Deterministic post-checks
# ---------------------------------------------------------------------------------------------
def _num(v):
    from services.documents import _num as n

    return n(v)


def iso6346_ok(cn: str) -> bool:
    cn = re.sub(r"[^A-Z0-9]", "", (cn or "").upper())
    if not re.fullmatch(r"[A-Z]{4}\d{7}", cn):
        return False
    values = {}
    v = 10
    for ch in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
        if v % 11 == 0:
            v += 1
        values[ch] = v
        v += 1
    total = sum((values[c] if c.isalpha() else int(c)) * (2 ** i) for i, c in enumerate(cn[:10]))
    return (total % 11) % 10 == int(cn[10])


def awb_ok(awb: str) -> bool:
    d = re.sub(r"\D", "", awb or "")
    if len(d) != 11:
        return False
    serial = int(d[3:10])
    return serial % 7 == int(d[10])


def _norm_value(kind: str, v):
    if v is None or v == "":
        return None
    if kind == "number":
        return _num(v)
    if kind in ("country", "currency", "incoterm"):
        return str(v).strip().upper()
    if kind == "list":
        if isinstance(v, list):
            return [str(x).strip() for x in v if str(x).strip()]
        return [x.strip() for x in re.split(r"[,;\n]+", str(v)) if x.strip()]
    if kind == "hs":
        d = re.sub(r"[^0-9]", "", str(v))
        return d or None
    if kind == "date":
        s = str(v).strip()[:10]
        try:
            date.fromisoformat(s)
            return s
        except ValueError:
            return str(v).strip()
    if kind in ("eori", "vat"):
        return re.sub(r"\s+", "", str(v)).upper()
    return str(v).strip() if not isinstance(v, (int, float)) else v


def _w(code, message, severity="warning", field=None):
    w = {"code": code, "severity": severity, "message": message, "field": field}
    i18n = warning_i18n(w)
    if i18n:
        w["message_i18n"] = i18n
    return w


# Finnish for every warning the reader writes: code -> (the English message as a regex, the Finnish template).
# The same table translates warnings stored before message_i18n existed (localize_extraction, at read time).
# The UI shows the field's own label next to the message, so the Finnish text does not repeat the field key.
_N = r"-?[\d,]+(?:\.\d+)?(?:e[+-]\d+)?"
WARN_FI = {
    "missing": (r"(?P<f>.+) was not found on the document", "Ei löytynyt asiakirjasta"),
    "country_code": (r"(?P<f>.+) '(?P<v>.*)' is not an ISO country code", "'{v}' ei ole ISO-maakoodi"),
    "currency": (r"Currency '(?P<v>.*)' is not a known ISO 4217 code", "Valuutta '{v}' ei ole tunnettu ISO 4217 -koodi"),
    "incoterm_old": (r"(?P<v>.+) is not an Incoterms 2020 rule; the closest 2020 rule is (?P<w>.+) \(RULE I-1\)\. "
                     r"Confirm with the seller\.",
                     "{v} ei ole Incoterms 2020 -toimituslauseke; lähin 2020-lauseke on {w} (sääntö I-1). Varmista myyjältä."),
    "incoterm": (r"'(?P<v>.*)' is not an Incoterms 2020 rule", "'{v}' ei ole Incoterms 2020 -toimituslauseke"),
    "eori_format": (r"EORI '(?P<v>.*)' does not match the EORI format \(RULE E-1\)",
                    "EORI '{v}' ei ole EORI-numeron muotoa (sääntö E-1)"),
    "eori_check": (r"EORI '(?P<v>.*)' has an invalid Finnish business ID check digit \(RULE E-2\)",
                   "EORI-numeron '{v}' Y-tunnuksen tarkistusnumero on väärä (sääntö E-2)"),
    "line_sum": (rf"Line items add up to (?P<n1>{_N}) but the subtotal is (?P<n2>{_N})",
                 "Rivit ovat yhteensä {n1}, mutta välisumma on {n2}"),
    "total_sum": (rf"Line items \((?P<n1>{_N})\) plus stated charges \((?P<n2>{_N})\) do not add up to the total (?P<n3>{_N})",
                  "Rivit ({n1}) ja ilmoitetut kulut ({n2}) eivät täsmää loppusummaan {n3}"),
    "line_math": (rf"Line (?P<i>\d+): (?P<n1>{_N}) × (?P<n2>{_N}) = (?P<n3>{_N}), but the line amount is (?P<n4>{_N})",
                  "Rivi {i}: {n1} × {n2} = {n3}, mutta rivin summa on {n4}"),
    "hs_format": (r"Line (?P<i>\d+): HS code '(?P<v>.*)' should be 6–10 digits", "Rivi {i}: HS-koodissa '{v}' pitää olla 6–10 numeroa"),
    "origin_missing": (r"Country of origin is not stated for every line; it is needed for the declaration \(DE 5/15\)",
                       "Alkuperämaata ei ole ilmoitettu joka rivillä; se tarvitaan tulli-ilmoitukseen (DE 5/15)"),
    "weights": (rf"Net weight (?P<n1>{_N}) kg is more than gross weight (?P<n2>{_N}) kg",
                "Nettopaino {n1} kg on suurempi kuin bruttopaino {n2} kg"),
    "container_check": (r"Container number (?P<v>\S+) fails the ISO 6346 check digit",
                        "Kontin numeron {v} ISO 6346 -tarkistusnumero ei täsmää"),
    "awb_check": (r"AWB (?P<v>\S+) fails the mod-7 check digit", "Lentorahtikirjan {v} tarkistusnumero (mod 7) ei täsmää"),
    "low_confidence": (r"(?P<f>.+) was hard to read \(confidence (?P<c>\d+)%\); check it against the document",
                       "Vaikea lukea (varmuus {c} %); tarkista arvo asiakirjasta"),
    "type_mismatch": (r"This looks like a (?P<a>.+), not a (?P<b>.+)\. Change the type if that is right\.",
                      "Tämä näyttää olevan {a}, ei {b}. Vaihda tyyppi, jos se pitää paikkansa."),
    "truncated": (r"Only the first (?P<n>\d+) pages were read; check later pages by hand",
                  "Vain {n} ensimmäistä sivua luettiin; tarkista loput sivut käsin"),
    "truncated_text": (r"Only the first part of the pages were read; check later pages by hand",
                       "Vain osa sivuista luettiin; tarkista loput sivut käsin"),
}
DOC_TYPE_FI = {"commercial invoice": "kauppalasku", "packing list": "pakkausluettelo", "bill of lading": "konossementti",
               "sea waybill": "merirahtikirja", "air waybill": "lentorahtikirja", "cmr": "CMR-rahtikirja",
               "cim": "CIM-rahtikirja", "certificate of origin": "alkuperätodistus", "eur1": "EUR.1",
               "dg declaration": "VAK-ilmoitus", "power of attorney": "valtakirja", "pod": "toimitustodiste",
               "customs decision": "tullipäätös", "other": "muu asiakirja"}


def _fi_num(s: str) -> str:
    """'25,392.85' -> '25 392,85' (Finnish digit grouping with a no-break space, decimal comma)."""
    return s.replace(",", "\u00a0").replace(".", ",")


def warning_i18n(w: dict) -> dict | None:
    """{en, fi} for a reader warning: the stored pair, or the Finnish built from its code and English message."""
    if (w.get("message_i18n") or {}).get("fi"):
        return w["message_i18n"]
    msg, code = w.get("message") or "", w.get("code")
    for key in ((code, "truncated_text") if code == "truncated" else (code,)):
        if key not in WARN_FI:
            continue
        m = re.fullmatch(WARN_FI[key][0], msg)
        if m:
            g = {k: (_fi_num(v) if re.fullmatch(r"n\d", k) else v) for k, v in m.groupdict().items()}
            if code == "type_mismatch":
                g = {k: DOC_TYPE_FI.get(v, v) for k, v in g.items()}
            return {"en": msg, "fi": WARN_FI[key][1].format(**g)}
    return None


def model_notes(raw: dict) -> list[dict]:
    """The model's notes as info warnings: {"en", "fi"} pairs become message + message_i18n (plain strings, from
    answers to the older prompt, stay English only)."""
    out = []
    for w in raw.get("warnings") or []:
        if isinstance(w, str) and w.strip():
            out.append(_w("model_note", w.strip()[:400], "info"))
        elif isinstance(w, dict) and isinstance(w.get("en"), str) and w["en"].strip():
            note = _w("model_note", w["en"].strip()[:400], "info")
            if isinstance(w.get("fi"), str) and w["fi"].strip():
                note["message_i18n"] = {"en": note["message"], "fi": w["fi"].strip()[:400]}
            out.append(note)
    return out


def localize_extraction(ex: dict | None) -> dict | None:
    """Add message_i18n to stored warnings (and reason_i18n to a reason built from them) for the Finnish UI."""
    if not ex:
        return ex
    ex = dict(ex)
    warnings = []
    for w in ex.get("warnings") or []:
        i18n = warning_i18n(w) if isinstance(w, dict) else None
        warnings.append({**w, "message_i18n": i18n} if i18n else w)
    if "warnings" in ex:
        ex["warnings"] = warnings
    reason = ex.get("reason")
    if reason and not ex.get("reason_i18n"):
        by_en = {w["message"]: w["message_i18n"]["fi"] for w in warnings if isinstance(w, dict) and w.get("message_i18n")}
        parts = reason.split("; ")
        if all(p in by_en for p in parts):
            ex["reason_i18n"] = {"en": reason, "fi": "; ".join(by_en[p] for p in parts)}
    return ex


# ISO 3166-1 alpha-2 (249 codes, from the IANA tz iso3166.tab) plus the EU customs geonomenclature extras:
# XK Kosovo, XS Serbia (old), XC Ceuta, XL Melilla, XI Northern Ireland, EU (origin "European Union").
COUNTRIES = frozenset("""
AD AE AF AG AI AL AM AO AQ AR AS AT AU AW AX AZ BA BB BD BE BF BG BH BI BJ BL BM BN BO BQ BR BS BT BV BW BY
BZ CA CC CD CF CG CH CI CK CL CM CN CO CR CU CV CW CX CY CZ DE DJ DK DM DO DZ EC EE EG EH ER ES ET FI FJ FK
FM FO FR GA GB GD GE GF GG GH GI GL GM GN GP GQ GR GS GT GU GW GY HK HM HN HR HT HU ID IE IL IM IN IO IQ IR
IS IT JE JM JO JP KE KG KH KI KM KN KP KR KW KY KZ LA LB LC LI LK LR LS LT LU LV LY MA MC MD ME MF MG MH MK
ML MM MN MO MP MQ MR MS MT MU MV MW MX MY MZ NA NC NE NF NG NI NL NO NP NR NU NZ OM PA PE PF PG PH PK PL PM
PN PR PS PT PW PY QA RE RO RS RU RW SA SB SC SD SE SG SH SI SJ SK SL SM SN SO SR SS ST SV SX SY SZ TC TD TF
TG TH TJ TK TL TM TN TO TR TT TV TW TZ UA UG UM US UY UZ VA VC VE VG VI VN VU WF WS YE YT ZA ZM ZW XK XS XC
XL XI EU
""".split())


def post_check(doc_type: str, fields: dict, items: list[dict]) -> list[dict]:
    from core.refs import ytunnus_valid

    warnings: list[dict] = []
    spec = FIELD_SPECS.get(doc_type) or FIELD_SPECS["other"]
    val = lambda k: (fields.get(k) or {}).get("value")  # noqa: E731
    for k, s in spec["fields"].items():
        if s["required"] and val(k) in (None, "", []):
            warnings.append(_w("missing", f"{k.replace('_', ' ').capitalize()} was not found on the document", "error", k))
    for k, f in fields.items():
        kind = (spec["fields"].get(k) or {}).get("kind")
        v = f.get("value")
        if v in (None, ""):
            continue
        if kind == "country" and str(v).upper() not in COUNTRIES:
            warnings.append(_w("country_code", f"{k.replace('_', ' ')} '{v}' is not an ISO country code", "warning", k))
            f["confidence"] = min(f.get("confidence") or 0, 0.5)
        if kind == "currency" and str(v) not in CURRENCIES:
            warnings.append(_w("currency", f"Currency '{v}' is not a known ISO 4217 code", "warning", k))
        if kind == "incoterm":
            if str(v) in OLD_INCOTERMS:
                warnings.append(_w("incoterm_old", f"{v} is not an Incoterms 2020 rule; the closest 2020 rule is "
                                                   f"{OLD_INCOTERMS[str(v)]} (RULE I-1). Confirm with the seller.", "warning", k))
            elif str(v) not in INCOTERMS_2020:
                warnings.append(_w("incoterm", f"'{v}' is not an Incoterms 2020 rule", "error", k))
        if kind == "eori":
            s = str(v)
            if not (re.fullmatch(r"FI\d{7}-\d", s) or re.fullmatch(r"FI\d{8}", s) or
                    (not s.startswith("FI") and re.fullmatch(r"[A-Z]{2}[A-Za-z0-9]{1,15}", s))):
                warnings.append(_w("eori_format", f"EORI '{s}' does not match the EORI format (RULE E-1)", "warning", k))
            elif s.startswith("FI") and not ytunnus_valid(f"{s[2:9]}-{s[-1]}"):
                warnings.append(_w("eori_check", f"EORI '{s}' has an invalid Finnish business ID check digit (RULE E-2)",
                                   "warning", k))
    if doc_type in ("commercial_invoice",):
        amounts = [_num(i.get("amount")) for i in items]
        if items and all(a is not None for a in amounts):
            s = round(sum(amounts), 2)
            sub, tot = _num(val("subtotal_amount")), _num(val("total_amount"))
            extras = sum(x for x in (_num(val("freight_amount")), _num(val("insurance_amount")),
                                     _num(val("other_charges_amount"))) if x)
            if sub is not None and abs(sub - s) > max(0.02, sub * 0.001):
                warnings.append(_w("line_sum", f"Line items add up to {s:,.2f} but the subtotal is {sub:,.2f}", "error",
                                   "subtotal_amount"))
            elif tot is not None and abs(tot - s - extras) > max(0.02, tot * 0.001) and abs(tot - s) > max(0.02, tot * 0.001):
                warnings.append(_w("total_sum", f"Line items ({s:,.2f}) plus stated charges ({extras:,.2f}) do not add up "
                                                f"to the total {tot:,.2f}", "error", "total_amount"))
        for i, it in enumerate(items):
            q, p, a = _num(it.get("quantity")), _num(it.get("unit_price")), _num(it.get("amount"))
            if q is not None and p is not None and a is not None and abs(q * p - a) > max(0.05, a * 0.005):
                warnings.append(_w("line_math", f"Line {i + 1}: {q:g} × {p:,.4g} = {q * p:,.2f}, but the line amount is {a:,.2f}",
                                   "warning", f"line_items.{i}"))
            hs = it.get("hs_code")
            if hs and not re.fullmatch(r"\d{6,10}", str(hs)):
                warnings.append(_w("hs_format", f"Line {i + 1}: HS code '{hs}' should be 6–10 digits", "warning", f"line_items.{i}"))
            if not it.get("country_of_origin") and not val("country_of_origin"):
                pass
        if items and not val("country_of_origin") and any(not it.get("country_of_origin") for it in items):
            warnings.append(_w("origin_missing", "Country of origin is not stated for every line; it is needed for the "
                                                 "declaration (DE 5/15)", "warning", "country_of_origin"))
    gw, nw = _num(val("total_gross_weight_kg")), _num(val("total_net_weight_kg"))
    if gw is not None and nw is not None and nw > gw:
        warnings.append(_w("weights", f"Net weight {nw:,.1f} kg is more than gross weight {gw:,.1f} kg", "error",
                           "total_net_weight_kg"))
    for k in ("container_numbers",):
        for c in val(k) or []:
            if not iso6346_ok(c):
                warnings.append(_w("container_check", f"Container number {c} fails the ISO 6346 check digit", "warning", k))
    if doc_type == "air_waybill" and val("awb_number") and not awb_ok(str(val("awb_number"))):
        warnings.append(_w("awb_check", f"AWB {val('awb_number')} fails the mod-7 check digit", "warning", "awb_number"))
    for k, f in fields.items():
        c = f.get("confidence")
        if c is not None and c < 0.7 and f.get("value") not in (None, ""):
            warnings.append(_w("low_confidence", f"{k.replace('_', ' ').capitalize()} was hard to read "
                                                 f"(confidence {c:.0%}); check it against the document", "info", k))
    return warnings


def normalise(doc_type: str, raw: dict) -> tuple[dict, list[dict]]:
    spec = FIELD_SPECS.get(doc_type) or FIELD_SPECS["other"]
    fields = {}
    for k, f in (raw.get("fields") or {}).items():
        if k not in spec["fields"]:
            continue
        if not isinstance(f, dict):
            f = {"value": f, "confidence": None}
        v = _norm_value(spec["fields"][k]["kind"], f.get("value"))
        if v in (None, "", []):
            continue
        conf = f.get("confidence")
        try:
            conf = max(0.0, min(1.0, float(conf))) if conf is not None else None
        except (TypeError, ValueError):
            conf = None
        fields[k] = {"value": v, "confidence": conf, "source": "ai"}
    items = []
    li_spec = spec.get("line_items") or {}
    for raw_it in raw.get("line_items") or []:
        if not isinstance(raw_it, dict):
            continue
        it = {}
        for k, s in li_spec.items():
            v = raw_it.get(k)
            if isinstance(v, dict) and "value" in v:
                v = v.get("value")
            v = _norm_value(s["kind"], v)
            if v not in (None, "", []):
                it[k] = v
        try:
            it["confidence"] = max(0.0, min(1.0, float(raw_it.get("confidence")))) if raw_it.get("confidence") is not None else None
        except (TypeError, ValueError):
            it["confidence"] = None
        if it.get("description") or it.get("amount") is not None:
            items.append(it)
    return fields, items


# ---------------------------------------------------------------------------------------------
# Job runner
# ---------------------------------------------------------------------------------------------
async def extract_document(doc_id: str) -> dict | None:
    doc = await db.documents.find_one({"id": doc_id}, NO_ID)
    if not doc:
        return None
    started = now_iso()
    ex = doc.get("extraction") or {}
    declared = doc.get("type") if doc.get("type_source") in ("user", "seed") else None
    result = {"status": "processing", "started_at": started, "attempts": int(ex.get("attempts") or 0)}
    try:
        data, _name, _ct = await get_file(doc["file_id"])
        inp = build_input(data, doc["content_type"])
    except Exception as exc:  # noqa: BLE001
        log.exception("could not read document %s", doc_id)
        result.update(status="failed", reason=f"The file could not be read ({type(exc).__name__}). Upload the original file.",
                      completed_at=now_iso())
        await _save(doc, result)
        return result
    if inp["method"] == "text" and len(inp["text"].strip()) < 20 and not inp["files"]:
        result.update(status="needs_review", method=inp["method"], completed_at=now_iso(),
                      reason="The document contains no readable text. Enter the fields by hand.")
        await _save(doc, result)
        return result
    from services.structured_docs import read_sheet
    raw = read_sheet(data, doc["content_type"], declared)  # an ERP export (CSV/XLSX) is read exactly, without a model
    if raw is None and not llm_available():
        result.update(status="needs_review", method=inp["method"], completed_at=now_iso(),
                      reason="AI is not configured on this server, so nothing was read automatically. "
                             "Enter the fields from the document.")
        await _save(doc, result)
        return result
    if raw is None:
        try:
            raw = await llm_json(AGENT, SYSTEM, build_prompt(declared, inp, doc.get("filename") or ""),
                                 files=inp["files"] or None, max_tokens=12000,
                                 ref={"type": "document", "id": doc_id}, company_id=doc.get("company_id"))
            if not isinstance(raw, dict):
                raise ValueError("Model returned no object")
        except LLMUnavailable as exc:
            result.update(status="needs_review", method=inp["method"], completed_at=now_iso(), reason=str(exc))
            await _save(doc, result)
            return result
        except Exception as exc:  # noqa: BLE001
            log.warning("extraction of %s failed: %s", doc_id, exc)
            result.update(status="needs_review", method=inp["method"], completed_at=now_iso(),
                          reason=f"The AI reader did not return a usable answer ({str(exc)[:160]}). "
                                 f"Retry, or enter the fields by hand.")
            await _save(doc, result)
            return result
    structured = raw.get("reader") == "structured"

    detected = raw.get("document_type") if raw.get("document_type") in DOC_TYPES else "other"
    tconf = raw.get("type_confidence")
    doc_type = declared or detected
    fields, items = normalise(doc_type, raw)
    if structured:
        for f in fields.values():
            f["source"] = "document"
    warnings = post_check(doc_type, fields, items)
    warnings += model_notes(raw)
    if declared and detected != declared and isinstance(tconf, (int, float)) and tconf >= 0.8:
        warnings.insert(0, _w("type_mismatch", f"This looks like a {detected.replace('_', ' ')}, not a "
                                               f"{declared.replace('_', ' ')}. Change the type if that is right.", "error"))
    if inp.get("truncated"):
        warnings.append(_w("truncated", f"Only the first {MAX_PAGES_VISION if inp['method'] == 'vision' else 'part of the'} "
                                        f"pages were read; check later pages by hand", "warning"))
    spec = FIELD_SPECS.get(doc_type) or FIELD_SPECS["other"]
    req = [k for k, s in spec["fields"].items() if s["required"]]
    req_conf = [(fields.get(k) or {}).get("confidence") or 0 for k in req]
    confs = [f["confidence"] for f in fields.values() if f.get("confidence") is not None]
    overall = round(min(req_conf) if req_conf else (sum(confs) / len(confs) if confs else 0), 3)
    blocking = [w for w in warnings if w["severity"] == "error"]
    status = "done" if not blocking and overall >= 0.8 and fields else "needs_review"
    result.update(
        status=status, method=inp["method"], model=None, detected_type=detected, detected_type_confidence=tconf,
        language=raw.get("language"), fields=fields, line_items=items, warnings=warnings, confidence=overall,
        completed_at=now_iso(), source="structured" if structured else "ai",
        reason=None if status == "done" else ("; ".join(w["message"] for w in blocking[:3]) or
                                              "Some values were hard to read; confirm them"),
    )
    run = None if structured else await db.ai_runs.find_one({"agent": AGENT, "ref.id": doc_id}, {"_id": 0, "model": 1, "id": 1},
                                                            sort=[("created_at", -1)])
    if structured:
        result["model"] = "rules/structured-sheet"
    if run:
        result["model"] = run.get("model")
        result["ai_run_id"] = run.get("id")
    changes = {}
    if doc.get("type_source") not in ("user", "seed"):
        changes = {"type": doc_type, "type_source": "ai"}
    await _save(doc, result, changes)
    return result


async def _save(doc: dict, result: dict, extra: dict | None = None):
    ex = {**(doc.get("extraction") or {}), **result}
    ex.pop("claimed_by", None)
    sets = {"extraction": ex, **(extra or {})}
    f = ex.get("fields") or {}
    for key in ("invoice_number", "bl_number", "awb_number", "cmr_number", "packing_list_number", "certificate_number"):
        if (f.get(key) or {}).get("value"):
            sets["number"] = f[key]["value"]
            break
    for key in ("invoice_date", "date_of_issue", "date", "issue_date"):
        if (f.get(key) or {}).get("value"):
            sets["issue_date"] = f[key]["value"]
            break
    await db.documents.update_one({"id": doc["id"]}, {"$set": sets})
    fresh = await db.documents.find_one({"id": doc["id"]}, NO_ID)
    if ex.get("status") in ("done", "needs_review", "failed"):
        await audit(agent_actor(AGENT), f"document_extraction_{ex['status']}", "document", doc["id"],
                    {"type": fresh.get("type"), "confidence": ex.get("confidence"), "reason": ex.get("reason")},
                    company_id=doc.get("company_id"))
        if fresh.get("shipment_id"):
            try:
                await run_tieout(fresh["shipment_id"])
                fresh = await db.documents.find_one({"id": doc["id"]}, NO_ID)
            except Exception:  # noqa: BLE001
                log.exception("tie-out failed for shipment %s", fresh.get("shipment_id"))
        if ex.get("status") in ("done", "needs_review"):
            await emit("document.extracted", {"document": fresh})


# ---------------------------------------------------------------------------------------------
# Durable queue worker
# ---------------------------------------------------------------------------------------------
_wake = None
_running = False
_tasks: set = set()


async def _claim() -> dict | None:
    from pymongo import ReturnDocument

    return await db.documents.find_one_and_update(
        {"extraction.status": "pending"},
        {"$set": {"extraction.status": "processing", "extraction.claimed_at": now_iso()},
         "$inc": {"extraction.attempts": 1}},
        sort=[("extraction.queued_at", 1)], projection=NO_ID, return_document=ReturnDocument.AFTER,
    )


async def drain(max_jobs: int = 50):
    global _running
    if _running:
        return
    _running = True
    try:
        for _ in range(max_jobs):
            job = await _claim()
            if not job:
                break
            try:
                await extract_document(job["id"])
            except Exception:  # noqa: BLE001
                log.exception("extraction job %s crashed", job["id"])
                await db.documents.update_one({"id": job["id"]}, {"$set": {
                    "extraction.status": "failed", "extraction.completed_at": now_iso(),
                    "extraction.reason": "Extraction crashed; retry or enter the fields by hand."}})
    finally:
        _running = False


def kick():
    """Start draining the queue now (called right after an upload)."""
    try:
        t = asyncio.get_running_loop().create_task(drain())
        _tasks.add(t)
        t.add_done_callback(_tasks.discard)
    except RuntimeError:
        pass


@background_loop(seconds=20)
async def extraction_worker():
    # Recover jobs orphaned by a restart: processing for > 10 minutes -> pending (max 3 attempts).
    from datetime import timedelta

    from core.util import iso

    cutoff = iso(now() - timedelta(minutes=10))
    await db.documents.update_many(
        {"extraction.status": "processing", "extraction.claimed_at": {"$lt": cutoff}, "extraction.attempts": {"$lt": 3}},
        {"$set": {"extraction.status": "pending"}})
    await db.documents.update_many(
        {"extraction.status": "processing", "extraction.claimed_at": {"$lt": cutoff}, "extraction.attempts": {"$gte": 3}},
        {"$set": {"extraction.status": "failed", "extraction.reason": "Extraction did not finish after 3 attempts."}})
    await drain()


async def requeue(doc_id: str):
    await db.documents.update_one({"id": doc_id}, {"$set": {
        "extraction.status": "pending", "extraction.queued_at": now_iso(), "extraction.attempts": 0,
        "extraction.reason": None}})
    kick()


def field_specs_public() -> dict:
    return {t: {"fields": [{"key": k, **v} for k, v in s["fields"].items()],
                "line_items": [{"key": k, **v} for k, v in (s.get("line_items") or {}).items()]}
            for t, s in FIELD_SPECS.items()}
