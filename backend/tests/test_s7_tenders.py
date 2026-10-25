"""s7 tenders against the real local backend: intake (xlsx / csv / JSON / e-mail), column mapping, pricing through the
real pricing engine (compared with POST /api/quotes), the filled workbook, return by e-mail, Autopilot and kill switch,
award. Fixtures are built with openpyxl in test_s7_tenders_support.py the way shippers send them.

Run: V7_API=http://localhost:8121/api V7_DB=huo_s7_tenders_1 DEMO_OPS_PASSWORD=ci-pass-7 pytest tests/test_s7_tenders.py
Optional TENDER_EVIDENCE_DIR: the filled workbook and its CSV export are copied there.
"""
from __future__ import annotations

import csv
import importlib.util
import io
import os
import time
import uuid

import openpyxl
import pytest

import test_s7_tenders_support as S

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")
HAS_TARIFFS = importlib.util.find_spec("services.contract_tariffs") is not None
EVIDENCE = os.environ.get("TENDER_EVIDENCE_DIR")


def upload(data: bytes, filename: str, *, company_id: str | None = None, price: bool = True) -> dict:
    files = {"file": (filename, data, "application/octet-stream")}
    form = {"company_id": company_id} if company_id else {}
    r = S.ops().post("/tenders", params={"price": "true"} if price else None, files=files, data=form)
    assert r.status_code == 200, r.text[:800]
    return r.json()["tender"]


def get(tid: str) -> dict:
    r = S.ops().get(f"/tenders/{tid}")
    assert r.status_code == 200, r.text[:300]
    return r.json()["tender"]


def wait_for(tid: str, cond, seconds: float = 90) -> dict:
    deadline = time.time() + seconds
    t = get(tid)
    while time.time() < deadline:
        if cond(t):
            return t
        time.sleep(1.5)
        t = get(tid)
    raise AssertionError(f"tender {tid} did not reach the state in {seconds}s: status={t['status']} agent={t.get('agent')}")


def by_ref(t: dict) -> dict:
    return {ln["ref"]: ln for ln in t["lanes"]}


@pytest.fixture(scope="module")
def company_id() -> str:
    return S.customer_user()["company_id"]


@pytest.fixture(scope="module")
def fi_tender(company_id) -> dict:
    return upload(S.finnish_workbook(), "Kuljetukset_2027_tarjouspyynto.xlsx", company_id=company_id)


@pytest.fixture()
def autopilot():
    """Turn Autopilot on for one test (typed confirmation, the real switch) and always back off."""
    def set_(on: bool):
        r = S.ops().post("/settings/autonomy/autopilot_enabled", json={"enabled": on, "confirmation": "enable"})
        assert r.status_code == 200, r.text[:300]
    set_(True)
    try:
        yield
    finally:
        S.mongo().settings.update_one({"key": "app"}, {"$set": {"autopilot_halted": False}})
        set_(False)


# ------------------------------------------------------------------------------------------------
# Parsing
# ------------------------------------------------------------------------------------------------
def test_finnish_sheet_header_and_column_mapping(fi_tender):
    t = fi_tender
    assert t["sheet"] == "Kuljetukset 2027" and t["header_rows"] == [5, 6] and t["data_start_row"] == 7
    fields = {k: v["field"] for k, v in t["columns_map"].items()}
    assert fields == {"A": "ref", "B": "origin_country", "C": "origin_zip", "D": "origin", "E": "destination_country",
                      "F": "destination_zip", "G": "destination", "H": "mode", "I": "equipment", "J": "weight_kg",
                      "K": "volume_m3", "L": "pallets", "M": "annual_volume", "N": "incoterm", "O": "commodity",
                      "P": "customer_price", "Q": "notes"}
    assert t["mapping_source"] == "synonym"
    # 12 lanes: title rows, the blank row, the total row and the footnote are not lanes
    assert [ln["ref"] for ln in t["lanes"]] == [str(i) for i in range(1, 13)]
    assert [ln["row"] for ln in t["lanes"]] == [7, 8, 9, 10, 11, 12, 14, 15, 16, 17, 18, 19]
    lanes = by_ref(t)
    assert lanes["2"]["parsed"]["weight_kg"] == 12000 and lanes["2"]["annual"]["volume"] == 104  # "12 000", "2/vk"
    assert lanes["4"]["parsed"]["weight_kg"] == 1500 and lanes["4"]["parsed"]["incoterm_place"] == "München"  # "1,5 t"
    assert lanes["6"]["parsed"]["containers"] == [{"type": "40HC", "count": 2}]  # "2 x 40HC"
    assert lanes["5"]["parsed"]["containers"] == [{"type": "40HC", "count": 1}]  # "40'HC"
    assert lanes["1"]["parsed"]["origin"]["zip"] == "01510" and lanes["1"]["parsed"]["origin"]["country"] == "FI"
    # the lane request is a POST /api/quotes body
    req = lanes["5"]["request"]
    assert req["origin"]["unlocode"] == "CNSHA" and req["destination"]["unlocode"] == "FIHEL" and req["mode"] == "sea"
    assert req["cargo"] == {"load_type": "FCL", "containers": [{"type": "40HC", "count": 1}], "weight_kg": 18000,
                            "description": "Huonekalut"}
    assert req["incoterm"] == "FOB" and req["incoterm_place"] == "Shanghai"
    assert lanes["2"]["request"]["destination"]["unlocode"] == "SEGOT"  # Göteborg = the network's Gothenburg
    assert lanes["1"]["request"]["cargo"]["packages"] == 6 and lanes["1"]["request"]["cargo"]["package_type"] == "pallet_eur"
    assert any(w["code"] == "incoterm.place_assumed" for w in lanes["1"]["warnings"])


def test_german_and_english_headers():
    from services import tenders as svc

    for data, name, expect in (
        (S.german_workbook(), "a.xlsx", {"origin", "origin_zip", "destination", "destination_zip", "mode", "weight_kg",
                                          "pallets", "annual_volume", "incoterm", "commodity", "ref"}),
        (S.english_csv(), "a.csv", {"ref", "origin", "origin_country", "destination", "destination_country", "mode",
                                     "equipment", "weight_kg", "volume_m3", "annual_volume", "incoterm"})):
        grids = svc.read_grids(data, svc.detect_format(data, name))
        table = svc.detect_table(grids)
        assert {m["field"] for m in table["mapping"].values() if m["field"]} == expect


# ------------------------------------------------------------------------------------------------
# Pricing
# ------------------------------------------------------------------------------------------------
def test_every_lane_priced_by_the_engine_or_explained(fi_tender):
    t = fi_tender
    assert t["status"] == "priced" and t["output_document_id"]
    lanes = by_ref(t)
    must_price = {"1": "road", "2": "road", "3": "road", "4": "road", "5": "sea", "6": "sea", "8": "sea", "9": "air", "10": "air"}
    for ref, mode in must_price.items():
        ln = lanes[ref]
        assert ln["status"] == "priced", (ref, ln["reason"])
        b = ln["best"]
        assert b["mode"] == mode and b["total_cents"] > 0 and b["carrier_name"] and b["valid_until"]
        assert b["annual_value_cents"] == round(b["total_cents"] * b["annual_shipments"])
        assert b["firmness"] in ("firm", "partial", "estimate", "indicative")
    assert lanes["5"]["best"]["unit"] == "per_container" and lanes["5"]["best"]["price_per_unit_cents"] == lanes["5"]["best"]["total_cents"]
    assert lanes["6"]["best"]["unit"] == "per_container"
    assert abs(lanes["6"]["best"]["price_per_unit_cents"] * 2 - lanes["6"]["best"]["total_cents"]) <= 1
    assert lanes["9"]["best"]["unit"] == "per_kg" and lanes["9"]["best"]["unit_qty"] >= 450
    assert lanes["2"]["best"]["annual_shipments"] == 104
    # not priced, each with its reason (never a made-up price)
    if t["engine"] == "pricing":
        assert lanes["7"]["status"] == "unpriced" and "No rates" in lanes["7"]["reason"] and lanes["7"]["best"] is None
    else:
        assert lanes["7"]["status"] == "priced" or lanes["7"]["reason"]
    assert lanes["11"]["status"] == "unpriced" and "Russia" in lanes["11"]["reason"] and "Venäjä" in lanes["11"]["reason_fi"]
    assert lanes["12"]["status"] == "unpriced" and "weight" in lanes["12"]["reason"] and "paino" in lanes["12"]["reason_fi"]
    tot = t["totals"]
    priced = [ln for ln in t["lanes"] if ln["status"] == "priced"]
    assert tot["lanes"] == 12 and tot["priced"] == len(priced) >= 9
    assert tot["annual_value_cents"] == sum(ln["best"]["annual_value_cents"] for ln in priced)
    assert tot["valid_until"] == min(ln["best"]["valid_until"] for ln in priced)


@pytest.mark.parametrize("ref", ["1", "5", "9"])
def test_lane_price_equals_post_quotes(fi_tender, company_id, ref):
    """The tender price is the price POST /api/quotes gives for the same request and customer."""
    ln = by_ref(fi_tender)[ref]
    r = S.ops().post("/quotes", json={**ln["request"], "company_id": company_id, "send": False})
    assert r.status_code == 200, r.text[:500]
    opts = [o for o in r.json()["quote"]["options"] if o["mode"] == ln["best"]["mode"]]
    same = [o for o in opts if o["id"] == ln["best"]["option_id"]]
    assert same, f"option {ln['best']['option_id']} not in the quote"
    assert same[0]["total_cents"] == ln["best"]["total_cents"]
    assert same[0]["carrier_name"] == ln["best"]["carrier_name"]
    if not any("firmness" in o for o in opts):  # s5 ranks firm before estimate; on main the cheapest wins
        assert ln["best"]["total_cents"] == min(o["total_cents"] for o in opts)


def test_reprice_keeps_lanes_and_bumps_revision(fi_tender):
    r = S.ops().post(f"/tenders/{fi_tender['id']}/price")
    assert r.status_code == 200, r.text[:300]
    t = r.json()["tender"]
    assert t["revision"] == fi_tender["revision"] + 1 and t["totals"]["priced"] == fi_tender["totals"]["priced"]
    assert t["output_document_id"] != fi_tender["output_document_id"]


# ------------------------------------------------------------------------------------------------
# Output workbook
# ------------------------------------------------------------------------------------------------
def _sheet_csv(ws) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    for row in ws.iter_rows(values_only=True):
        w.writerow(["" if v is None else v for v in row])
    return buf.getvalue()


def test_filled_workbook_keeps_the_customer_sheet_and_adds_prices(fi_tender):
    t = get(fi_tender["id"])
    r = S.ops().get(f"/tenders/{t['id']}/file")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:2] == ["Ohjeet", "Kuljetukset 2027"] and wb.sheetnames[2].startswith("Yhteenveto")
    ws = wb["Kuljetukset 2027"]
    # the customer's sheet as sent: titles, merged groups, colours, formula, footnote
    assert ws["A1"].value == "Esimerkki Teollisuus Oy – Rahtikilpailutus 2027"
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"A1:Q1", "B5:D5", "E5:G5", "A5:A6"} <= merged
    assert ws["A5"].fill.fgColor.rgb.endswith("1F3864") and ws.freeze_panes == "B7"
    assert ws["M20"].value == "=SUM(M7:M19)" and ws["A22"].value.startswith("* Hinnat")
    for c, v in enumerate(S.FI_HEADER_SUB, start=1):
        assert ws.cell(6, c).value == v
    # price columns right of the header (Q is the last customer column)
    assert ws["R6"].value == "Kuljetusmuoto / Mode" and ws["X6"].value.startswith("Hinta / lähetys EUR")
    assert ws["AC6"].value == "Huomautukset / Notes"
    headers = {ws.cell(6, c).value: c for c in range(18, 30)}
    for ln in t["lanes"]:
        assert ws.cell(ln["row"], 1).value == int(ln["ref"])  # original row untouched
        per_ship = ws.cell(ln["row"], headers["Hinta / lähetys EUR / Price per shipment EUR"]).value
        notes = ws.cell(ln["row"], headers["Huomautukset / Notes"]).value
        if ln["status"] == "priced":
            assert per_ship == ln["best"]["total_cents"] / 100
            annual = ws.cell(ln["row"], headers["Vuosiarvo EUR / Annual value EUR"]).value
            assert annual == ln["best"]["annual_value_cents"] / 100
        else:
            assert per_ship is None and notes and ln["reason"] in notes
    summ = wb[wb.sheetnames[2]]
    rows = {row[0]: row[1] for row in summ.iter_rows(values_only=True) if row[0]}
    assert rows["Vuosiarvo yhteensä EUR / Total annual value EUR"] == t["totals"]["annual_value_cents"] / 100
    assert rows["Reittejä / Lanes"] == 12 and rows["Hinnat voimassa / Prices valid until"] == t["totals"]["valid_until"]
    settings = S.mongo().settings.find_one({"key": "app"}) or {}
    from core.settings import DEFAULTS
    assert rows["Ehdot / Terms"] == settings.get("terms_version", DEFAULTS["terms_version"])
    if EVIDENCE:
        os.makedirs(EVIDENCE, exist_ok=True)
        with open(os.path.join(EVIDENCE, "filled_tender_fi.xlsx"), "wb") as fh:
            fh.write(r.content)
        with open(os.path.join(EVIDENCE, "filled_tender_fi.csv"), "w", encoding="utf-8-sig", newline="") as fh:
            fh.write(_sheet_csv(ws))
        with open(os.path.join(EVIDENCE, "filled_tender_fi_summary.csv"), "w", encoding="utf-8-sig", newline="") as fh:
            fh.write(_sheet_csv(summ))


def test_csv_in_xlsx_out():
    t = upload(S.english_csv(), "lanes_2027.csv")
    assert t["original_format"] == "csv" and t["company_id"] is None
    lanes = by_ref(t)
    assert lanes["L-01"]["status"] == "priced" and lanes["L-01"]["best"]["mode"] == "road"
    assert lanes["L-02"]["status"] == "priced" and lanes["L-02"]["best"]["unit"] == "per_container"
    assert lanes["L-03"]["status"] == "priced" and lanes["L-03"]["best"]["mode"] == "air"
    assert lanes["L-04"]["annual"]["volume"] == 52  # "weekly"
    r = S.ops().get(f"/tenders/{t['id']}/file")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.worksheets[0]
    assert ws["A1"].value == "Lane ID" and ws["A2"].value == "L-01" and ws["H2"].value == 5400
    assert ws["L1"].value == "Kuljetusmuoto / Mode" and ws["R2"].value == lanes["L-01"]["best"]["total_cents"] / 100


def test_json_lanes():
    body = {"name": "API tender", "price": True, "lanes": [
        {"ref": "A", "origin": "FIHEL", "destination": "SESTO", "mode": "road", "weight_kg": 1000, "pallets": 2,
         "annual_volume": 10},
        {"ref": "B", "request": {"origin": "CNSHA", "destination": "FIHEL", "mode": "sea", "incoterm": "FOB",
                                 "incoterm_place": "Shanghai", "cargo": {"containers": [{"type": "20GP", "count": 1}]}},
         "annual_volume": 4}]}
    r = S.ops().post("/tenders", json=body)
    assert r.status_code == 200, r.text[:500]
    t = r.json()["tender"]
    assert t["source"] == "api" and t["status"] == "priced"
    assert all(ln["status"] == "priced" and ln["best"]["total_cents"] > 0 for ln in t["lanes"])
    assert by_ref(t)["B"]["best"]["annual_value_cents"] == by_ref(t)["B"]["best"]["total_cents"] * 4
    f = S.ops().get(f"/tenders/{t['id']}/file")
    assert f.status_code == 200 and openpyxl.load_workbook(io.BytesIO(f.content)).worksheets[0]["A1"].value == "ref"


def test_bad_files_are_refused_plainly():
    r = S.ops().post("/tenders", files={"file": ("x.pdf", b"%PDF-1.4 \x00\x01\x02", "application/pdf")})
    assert r.status_code == 415
    r = S.ops().post("/tenders", files={"file": ("notes.csv", "a;b\n1;2\n".encode(), "text/csv")})
    assert r.status_code == 422 and "origin" in r.text


def test_unknown_column_is_reported_not_guessed():
    from core.llm import llm_available
    data = (S.CSV_TEXT.replace("Incoterm\n", "Incoterm;Määrä\n", 1).replace("FCA Hamburg\n", "FCA Hamburg;3\n", 1)).encode()
    t = upload(data, "lanes_extra.csv", price=False)
    entry = next(v for v in t["columns_map"].values() if v["header"] == "Määrä")
    if not llm_available():
        assert entry["field"] is None and "Määrä" in t["unmapped_columns"]
        assert any(w["code"] == "columns.unmapped" for w in t["warnings"])
    else:  # live model: only an allowed field, never a price
        assert entry["field"] in (None, *[f for f in __import__("services.tenders", fromlist=["FIELDS"]).FIELDS if f != "customer_price"])


# ------------------------------------------------------------------------------------------------
# E-mail intake, return, Autopilot, kill switch
# ------------------------------------------------------------------------------------------------
def post_eml(raw: bytes) -> list[dict]:
    r = S.ops().post("/tenders/inbound-email", content=raw, headers={"Content-Type": "message/rfc822"})
    assert r.status_code == 200, r.text[:500]
    return r.json()["tenders"]


def outbox(tid: str) -> list[dict]:
    return list(S.mongo().email_outbox.find({"ref.id": tid, "tag": "tender"}, {"_id": 0}))


def test_email_tender_waits_for_ops_when_autopilot_is_off(company_id):
    S.mongo().settings.update_one({"key": "app"}, {"$set": {"autopilot_enabled": False, "autopilot_halted": False}})
    raw = S.rfq_email("customer.demo@huolitsija.fi")
    items = post_eml(raw)
    assert len(items) == 1
    t0 = items[0]
    assert t0["source"] == "email" and t0["company_id"] == company_id and t0["requester_email"] == "customer.demo@huolitsija.fi"
    assert {m["field"] for m in t0["columns_map"].values() if m["field"]} >= {"origin", "destination", "weight_kg", "annual_volume"}
    t = wait_for(t0["id"], lambda x: x["status"] == "priced" and x["agent"]["state"] == "owner_queue")
    assert {r["code"] for r in t["agent"]["reasons"]} >= {"autopilot_off"} or t["agent"].get("decision_id")
    assert all(ln["status"] == "priced" for ln in t["lanes"]), [ln["reason"] for ln in t["lanes"]]
    assert outbox(t["id"]) == []
    # the same message again is the same tender
    assert [x["id"] for x in post_eml(raw)] == [t["id"]]
    # ops returns it
    r = S.ops().post(f"/tenders/{t['id']}/return", json={})
    assert r.status_code == 200, r.text[:300]
    t = r.json()["tender"]
    assert t["status"] == "returned" and t["returned_to"] == "customer.demo@huolitsija.fi"
    mails = outbox(t["id"])
    assert len(mails) == 1 and t["number"] in mails[0]["subject"] and f"/api/tenders/{t['id']}/file?t=" in mails[0]["text"]
    # the filled workbook itself is attached (a dict passed where core.email expects a tuple once sent a 12-byte "filename")
    att = mails[0].get("attachments") or []
    assert len(att) == 1 and att[0]["filename"].endswith(".xlsx"), att
    assert att[0]["content_type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" and att[0]["size"] > 1000, att
    link = mails[0]["text"].split("Lataa hinnoiteltu taulukko / Download the priced sheet: ")[1].split()[0]
    token = link.split("?t=")[1]
    anon = __import__("httpx").get(f"{S.API}/tenders/{t['id']}/file", params={"t": token}, timeout=30)
    assert anon.status_code == 200 and anon.content[:2] == b"PK"
    assert __import__("httpx").get(f"{S.API}/tenders/{t['id']}/file", params={"t": token + "x"}, timeout=30).status_code == 401
    assert __import__("httpx").get(f"{S.API}/tenders/{t['id']}/file", timeout=30).status_code == 401


def test_email_tender_is_returned_without_a_person_when_autopilot_is_on(autopilot):
    items = post_eml(S.rfq_email("customer.demo@huolitsija.fi", subject="Tarjouspyyntö maantiekuljetukset 2027"))
    t = wait_for(items[0]["id"], lambda x: x["status"] == "returned", 120)
    assert t["agent"]["state"] == "returned" and t["returned_to"] == "customer.demo@huolitsija.fi"
    assert len(outbox(t["id"])) == 1
    act = S.mongo().autonomy_actions.find_one({"entity_id": t["id"], "kind": "tender_returned"}) or \
        S.mongo().agent_decisions.find_one({"subject_id": t["id"], "action": "return_tender"})
    assert act, "no autonomy record for the return"


def test_kill_switch_holds_the_agent(autopilot):
    S.mongo().settings.update_one({"key": "app"}, {"$set": {"autopilot_halted": True}})
    items = post_eml(S.rfq_email("customer.demo@huolitsija.fi", subject="RFQ 2027 lanes (halt test)"))
    # held by our own gate ("held") or, with s2's core.jobs, the agent job is never claimed while halted ("pending")
    t = wait_for(items[0]["id"], lambda x: x["agent"]["state"] in ("held", "pending"), 30)
    time.sleep(4)
    t = get(t["id"])
    assert t["status"] == "parsed" and outbox(t["id"]) == []
    S.mongo().settings.update_one({"key": "app"}, {"$set": {"autopilot_halted": False}})
    t = wait_for(t["id"], lambda x: x["status"] == "returned", 150)  # resumed by the sweep / job worker
    assert len(outbox(t["id"])) == 1


def test_email_without_tender_sheet_creates_nothing():
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Tuote", "Hinta"])
    wb.active.append(["Tuoli", 12])
    buf = io.BytesIO()
    wb.save(buf)
    assert post_eml(S.rfq_email(f"x{uuid.uuid4().hex[:6]}@example.com", attachment=buf.getvalue(), filename="hinnasto.xlsx")) == []


# ------------------------------------------------------------------------------------------------
# Award / lost
# ------------------------------------------------------------------------------------------------
@pytest.mark.skipif(HAS_TARIFFS, reason="services.contract_tariffs is deployed")
def test_award_without_contract_tariffs_says_so(fi_tender):
    r = S.ops().post(f"/tenders/{fi_tender['id']}/award", json={})
    assert r.status_code == 501 and "contract_tariffs" in r.text


@pytest.mark.skipif(not HAS_TARIFFS, reason="services.contract_tariffs (another s7 builder) is not on this branch yet")
def test_award_creates_contract_tariff(company_id):
    t = upload(S.finnish_workbook(), "award.xlsx", company_id=company_id)
    lane_ids = [ln["id"] for ln in t["lanes"] if ln["status"] == "priced"][:3]
    r = S.ops().post(f"/tenders/{t['id']}/award", json={"lane_ids": lane_ids})
    assert r.status_code == 200, r.text[:500]
    body = r.json()
    assert body["tender"]["status"] == "awarded" and body["tariff"]["id"] == body["tender"]["award"]["tariff_id"]


def test_lost_and_listing():
    t = upload(S.english_csv(), "lost.csv", price=False)
    assert t["status"] == "parsed"
    r = S.ops().post(f"/tenders/{t['id']}/lost", json={"reason": "Asiakas valitsi toisen / customer chose another"})
    assert r.status_code == 200 and r.json()["tender"]["status"] == "lost"
    lst = S.ops().get("/tenders", params={"status": "lost"}).json()
    assert t["id"] in [x["id"] for x in lst["items"]] and lst["counts"]["lost"] >= 1
    assert S.ops().post(f"/tenders/{t['id']}/price").status_code == 409
