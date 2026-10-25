"""Helpers for test_s7_tenders.py (no tests here): API clients and the tender fixtures, built with openpyxl the way
customers send them (title rows, merged group headers, Finnish / English / German column names, blank and total rows).

Env: V7_API (http://localhost:8120/api), V7_DB, MONGO_URL, DEMO_OPS_PASSWORD.
`python tests/test_s7_tenders_support.py` writes the fixtures to tests/fixtures/s7_tender_*.
"""
from __future__ import annotations

import io
import os
import sys
from email.message import EmailMessage
from email.utils import make_msgid
from functools import lru_cache

import httpx
import pymongo

API = os.environ.get("V7_API", "http://localhost:8120/api")
DB_NAME = os.environ.get("V7_DB", "huo_s7_tenders_1")
OPS_PW = os.environ.get("DEMO_OPS_PASSWORD", "ci-pass-7")
os.environ.setdefault("DB_NAME", DB_NAME)  # before any backend module reads core.config
BACKEND = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)


@lru_cache
def mongo():
    return pymongo.MongoClient(os.environ.get("MONGO_URL", "mongodb://localhost:27017"))[DB_NAME]


def login(email: str, password: str) -> httpx.Client:
    r = httpx.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login {email}: {r.status_code} {r.text[:300]}"
    return httpx.Client(base_url=API, timeout=180, headers={"Authorization": f"Bearer {r.json()['token']}",
                                                            "X-Include-Test-Tenants": "true"})


@lru_cache
def ops() -> httpx.Client:
    return login("ops.demo@huolitsija.fi", OPS_PW)


def customer_user() -> dict:
    u = mongo().users.find_one({"email": "customer.demo@huolitsija.fi"}) or mongo().users.find_one({"org_type": "customer"})
    assert u and u.get("company_id"), "no demo customer user in the seeded DB"
    return u


# ------------------------------------------------------------------------------------------------
# Fixtures
# ------------------------------------------------------------------------------------------------
# Finnish sheet: 3 title rows, a blank row, a two-row header with merged "Lähtö" / "Määränpää" groups.
FI_HEADER_TOP = ["Nro", "Lähtö", None, None, "Määränpää", None, None, "Kuljetusmuoto", "Kontti", "Paino kg",
                 "Tilavuus m3", "Lavat", "Määrä vuodessa", "Incoterm", "Tavara", "Hinta EUR / yksikkö", "Huomautukset"]
FI_HEADER_SUB = [None, "Maa", "Postinumero", "Kaupunki", "Maa", "Postinumero", "Kaupunki"]
FI_LANES = [
    # nro, lähtö maa, postinro, kaupunki, määränpää maa, postinro, kaupunki, muoto, kontti, paino, m3, lavat, vuodessa, incoterm, tavara, hinta, huom
    [1, "FI", "01510", "Vantaa", "SE", "11122", "Stockholm", "Maantie", None, 3200, 8.5, 6, 52, "DAP", "Paperirullat", None, None],
    [2, "FI", "33100", "Tampere", "SE", "41101", "Göteborg", "Maantie", None, "12 000", 30, 24, "2/vk", "DAP", "Koneenosat", None, None],
    [3, "FI", "00980", "Helsinki", "DE", "20457", "Hamburg", "Maantie", None, 800, 2.4, 2, 24, None, "Elektroniikka", None, None],
    [4, "FI", "20100", "Turku", "DE", "80331", "München", "Maantie", None, "1,5 t", 6, 4, 12, "DAP München", "Työkalut", None, None],
    [5, "CN", None, "Shanghai", "FI", "00980", "Helsinki", "Meri", "40'HC", 18000, None, None, 26, "FOB Shanghai", "Huonekalut", None, None],
    [6, "CN", None, "Ningbo", "FI", "48310", "Kotka", "Meri", "2 x 40HC", 36000, None, None, 12, "FOB", "Kodinkoneet", None, None],
    None,  # blank row inside the table
    [7, "FI", "00980", "Helsinki", "CN", None, "Shanghai", "Meri", "40HC", 20000, None, None, 10, "CIF", "Sellu", None, None],
    [8, "FI", "48310", "Kotka", "US", None, "New York", "Meri", "40 HC", 15000, None, None, 8, None, "Saha-tavara", None, None],
    [9, "FI", "01530", "Helsinki", "US", "10001", "New York", "Lento", None, 450, 2.4, 2, 24, None, "Varaosat", None, "kiire"],
    [10, "FI", "01530", "Helsinki", "US", "60666", "Chicago", "Lento", None, 1200, 6, 4, 6, None, "Mittalaitteet", None, None],
    [11, "FI", "00980", "Helsinki", "RU", None, "Pietari", "Maantie", None, 5000, 20, 12, 12, None, "Elintarvikkeet", None, None],
    [12, "FI", "90100", "Oulu", "SE", None, "Haparanda", "Maantie", None, None, None, 3, 20, None, "Kalusteet", None, "paino myöhemmin"],
]


def finnish_workbook() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    info = wb.active
    info.title = "Ohjeet"
    info["A1"] = "Täytä hinnat välilehdelle Kuljetukset 2027. Vastaukset 15.11.2026 mennessä."
    ws = wb.create_sheet("Kuljetukset 2027")
    width = len(FI_HEADER_TOP)
    ws.cell(1, 1, "Esimerkki Teollisuus Oy – Rahtikilpailutus 2027").font = Font(bold=True, size=16)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.cell(2, 1, "Tarjouspyyntö / Request for quotation · vastaus 15.11.2026 mennessä")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.cell(3, 1, "Hinnat EUR alv 0 %, voimassa 1.1.–31.12.2027 · yhteyshenkilö Matti Meikäläinen")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width)
    thin = Side(style="thin", color="999999")
    head_fill = PatternFill("solid", fgColor="1F3864")
    for c, v in enumerate(FI_HEADER_TOP, start=1):
        if v is not None:
            cell = ws.cell(5, c, v)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c, v in enumerate(FI_HEADER_SUB, start=1):
        if v is not None:
            cell = ws.cell(6, c, v)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
    ws.merge_cells("B5:D5")
    ws.merge_cells("E5:G5")
    for col in ("A", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"):
        ws.merge_cells(f"{col}5:{col}6")
    r = 7
    for lane in FI_LANES:
        if lane is not None:
            for c, v in enumerate(lane, start=1):
                if v is not None:
                    ws.cell(r, c, v)
        r += 1
    ws.cell(r, 1, "Yhteensä")
    ws.cell(r, 13, "=SUM(M7:M19)")
    ws.cell(r + 2, 1, "* Hinnat sisältäen polttoainelisän; lisät eriteltävä erikseen.")
    ws.freeze_panes = "B7"
    for col, w in (("A", 6), ("D", 16), ("G", 16), ("O", 18), ("P", 18), ("Q", 20)):
        ws.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


CSV_TEXT = (
    "Lane ID;Origin;Origin country;Destination;Destination country;Mode;Equipment;Weight (kg);Volume (cbm);"
    "Shipments per year;Incoterm\n"
    "L-01;Hamburg;DE;Helsinki;FI;Road;;5400;18;48;FCA Hamburg\n"
    "L-02;Shanghai;CN;Helsinki;FI;Sea;40HC;;;20;FOB Shanghai\n"
    "L-03;Helsinki;FI;New York;US;Air;;300;1,5;12;\n"
    "L-04;Tallinn;EE;Helsinki;FI;Road;;2000;6;weekly;DAP Helsinki\n"
)


def english_csv() -> bytes:
    return CSV_TEXT.encode("utf-8")


DE_ROWS = [["Ausschreibung Landverkehr 2027"], [],
           ["Pos.", "Von", "Von PLZ", "Nach", "Nach PLZ", "Verkehrsträger", "Gewicht (kg)", "Paletten", "Sendungen pro Jahr",
            "Lieferbedingung", "Ware"],
           [1, "Hamburg", "20457", "Helsinki", "00980", "LKW", 6400, 10, 40, "FCA Hamburg", "Ersatzteile"],
           [2, "Duisburg", "47051", "Helsinki", "00980", "LKW", 2100, 4, 26, "FCA Duisburg", "Maschinenteile"],
           [3, "Stockholm", "11122", "Turku", "20100", "LKW", 900, 2, 52, "FCA Stockholm", "Papier"]]


def german_workbook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Lanes"
    for r, row in enumerate(DE_ROWS, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rfq_email(sender: str, *, subject: str = "Ausschreibung / Tarjouspyyntö 2027", attachment: bytes | None = None,
              filename: str = "Ausschreibung_2027.xlsx") -> bytes:
    msg = EmailMessage()
    msg["From"] = f"Einkauf <{sender}>"
    msg["To"] = "myynti@huolitsija.fi"
    msg["Subject"] = subject
    msg["Message-ID"] = make_msgid(domain="example-shipper.de")
    msg.set_content("Hei,\n\nliitteenä vuoden 2027 kuljetusten tarjouspyyntö. Pyydämme hinnat taulukkoon.\n\n"
                    "Guten Tag, anbei unsere Ausschreibung für 2027.\n\nTerveisin / Mit freundlichen Grüßen\nEinkauf")
    msg.add_attachment(attachment if attachment is not None else german_workbook(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=filename)
    return bytes(msg)


if __name__ == "__main__":
    here = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
    os.makedirs(here, exist_ok=True)
    for name, data in (("s7_tender_fi_kuljetukset_2027.xlsx", finnish_workbook()), ("s7_tender_en_lanes.csv", english_csv()),
                       ("s7_tender_de_ausschreibung.xlsx", german_workbook()),
                       ("s7_tender_rfq_email.eml", rfq_email("customer.demo@huolitsija.fi"))):
        with open(os.path.join(here, name), "wb") as fh:
            fh.write(data)
        print("wrote", name, len(data))
